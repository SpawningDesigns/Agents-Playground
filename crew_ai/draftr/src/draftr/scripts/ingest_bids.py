#!/usr/bin/env python
"""
Knowledge Vault ingestion pipeline.

Reads past winning bid documents from the configured bids directory,
chunks them, embeds each chunk with a local sentence-transformer model,
and upserts everything into a ChromaDB HTTP collection.

Supported formats: .pdf  .txt  .md  .docx

All defaults are read from config.yaml. CLI flags override config values.

Usage (via installed CLI)
-------------------------
    ingest_bids
    ingest_bids --reset
    ingest_bids --bids-dir /path/to/bids --chunk-size 800

Usage (direct)
--------------
    python scripts/ingest_bids.py
"""

import argparse
import hashlib
import logging
import sys
from pathlib import Path
from typing import Optional

import chromadb

from draftr.vault_config import DEFAULT_CONFIG, SentenceTransformerEF, VaultConfig, load_vault_config

log = logging.getLogger(__name__)

SUPPORTED_FORMATS = {".pdf", ".txt", ".md", ".docx"}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def _extract_pdf(path: Path) -> list[tuple[str, int]]:
    """Return (text, page_number) for every non-empty PDF page."""
    try:
        from pypdf import PdfReader
    except ImportError:
        log.error("pypdf is required for PDF files: pip install pypdf")
        return []

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((text, i))
    return pages


def _extract_docx(path: Path) -> list[tuple[str, int]]:
    """Return the full document text as a single (text, page=1) entry."""
    try:
        from docx import Document
    except ImportError:
        log.error("python-docx is required for .docx files: pip install python-docx")
        return []

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    full_text = "\n".join(paragraphs)
    return [(full_text, 1)] if full_text.strip() else []


def _extract_text(path: Path) -> list[tuple[str, int]]:
    """Return the whole plain-text file as a single (text, page=1) entry."""
    try:
        return [(path.read_text(encoding="utf-8", errors="replace"), 1)]
    except Exception as exc:
        log.warning("Could not read %s: %s", path.name, exc)
        return []


def extract(path: Path) -> list[tuple[str, int]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    return _extract_text(path)


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

def chunk(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Split text into overlapping fixed-size character windows."""
    text = " ".join(text.split())   # normalise whitespace
    results, start = [], 0
    while start < len(text):
        piece = text[start: start + chunk_size]
        if len(piece.strip()) > 50:
            results.append(piece)
        start += chunk_size - overlap
    return results


# ---------------------------------------------------------------------------
# Stable document ID
# ---------------------------------------------------------------------------

def doc_id(source_file: str, global_chunk_index: int) -> str:
    raw = f"{source_file}::{global_chunk_index}"
    return hashlib.sha256(raw.encode()).hexdigest()[:16]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main(
    bids_dir: Optional[Path] = None,
    chunk_size: Optional[int] = None,
    overlap: Optional[int] = None,
    reset: bool = False,
    config_path: Path = DEFAULT_CONFIG,
) -> None:
    """
    Ingest bid documents into ChromaDB.
    Positional args override config.yaml; config.yaml overrides built-in defaults.
    """
    vc: VaultConfig = load_vault_config(config_path)

    effective_bids_dir = bids_dir or vc.bids_dir
    effective_chunk_size = chunk_size if chunk_size is not None else vc.chunk_size
    effective_overlap = overlap if overlap is not None else vc.overlap

    if not effective_bids_dir.exists():
        log.error("Bids directory not found: %s — create it and drop your bid documents inside.", effective_bids_dir)
        sys.exit(1)

    log.info("Connecting to ChromaDB at %s:%s", vc.host, vc.port)
    try:
        client = chromadb.HttpClient(host=vc.host, port=vc.port)
        client.heartbeat()
    except Exception as exc:
        log.error("Cannot reach ChromaDB at %s:%s — %s", vc.host, vc.port, exc)
        log.error("Make sure the ChromaDB server is running.")
        sys.exit(1)

    log.info("Embedding model : %s", vc.embedding_model)
    ef = SentenceTransformerEF(model_name=vc.embedding_model)
    log.info("Collection      : %s", vc.collection)
    log.info("Chunk size      : %d  overlap: %d", effective_chunk_size, effective_overlap)

    if reset:
        try:
            client.delete_collection(vc.collection)
            log.info("Collection '%s' deleted — rebuilding from scratch.", vc.collection)
        except Exception:
            pass

    collection = client.get_or_create_collection(
        name=vc.collection,
        embedding_function=ef,
        metadata={"hnsw:space": "cosine"},
    )

    files = sorted(
        f for f in effective_bids_dir.rglob("*") if f.suffix.lower() in SUPPORTED_FORMATS
    )
    if not files:
        log.warning("No supported files found in %s — add .pdf/.txt/.md/.docx files and re-run.", effective_bids_dir)
        sys.exit(0)

    log.info("Found %d file(s) in %s", len(files), effective_bids_dir)

    global_chunk_idx = 0
    total_files_ok = 0

    for file in files:
        log.info("Processing: %s", file.name)
        pages = extract(file)
        if not pages:
            log.warning("  Skipping %s — no extractable content.", file.name)
            continue

        file_chunks = 0
        for text, page_num in pages:
            pieces = chunk(text, effective_chunk_size, effective_overlap)
            ids, docs, metas = [], [], []
            for piece in pieces:
                ids.append(doc_id(file.name, global_chunk_idx))
                docs.append(piece)
                metas.append(
                    {
                        "source_file": file.name,
                        "page_number": page_num,
                        "chunk_index": global_chunk_idx,
                    }
                )
                global_chunk_idx += 1
                file_chunks += 1

            if ids:
                collection.upsert(ids=ids, documents=docs, metadatas=metas)

        log.info("  → %d chunk(s) ingested", file_chunks)
        total_files_ok += 1

    log.info(
        "Done — %d total chunks from %d file(s) | ChromaDB %s:%s | collection: %s | total in collection: %d",
        global_chunk_idx, total_files_ok, vc.host, vc.port, vc.collection, collection.count(),
    )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Ingest past winning bid documents into the Knowledge Vault. "
            "Defaults are read from config.yaml; these flags override them."
        )
    )
    parser.add_argument(
        "--bids-dir", metavar="DIR",
        help="Directory containing bid documents (overrides ingestion.bids_dir in config.yaml)",
    )
    parser.add_argument(
        "--chunk-size", type=int, metavar="N",
        help="Characters per chunk (overrides ingestion.chunk_size in config.yaml)",
    )
    parser.add_argument(
        "--overlap", type=int, metavar="N",
        help="Overlap between chunks (overrides ingestion.overlap in config.yaml)",
    )
    parser.add_argument(
        "--reset", action="store_true",
        help="Delete and recreate the collection before ingesting",
    )
    parser.add_argument(
        "--config", metavar="FILE", default=str(DEFAULT_CONFIG),
        help=f"Path to config.yaml (default: {DEFAULT_CONFIG})",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    # ingest_bids is run standalone, so it sets up logging itself
    from draftr.main import setup_logging
    setup_logging(Path(args.config))
    main(
        bids_dir=Path(args.bids_dir) if args.bids_dir else None,
        chunk_size=args.chunk_size,
        overlap=args.overlap,
        reset=args.reset,
        config_path=Path(args.config),
    )
